import openpyxl, xlrd, json, re
BANKS = {'HCB': 'Home Credit', 'HSBK': 'Народный', 'KSPI': 'Kaspi Bank', 'CCBN': 'ЦентрКредит',
         'ASBN': 'ForteBank', 'FRHC': 'Фридом'}  # HCB первым: с 05.2026 в названии есть «ForteBank»

def match(name):
    for k, v in BANKS.items():
        if name and v in str(name):
            return k

def sheets_xlsx(fn):
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    return {re.sub(r'\s+', ' ', s.strip()): [list(r) for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames}

def sheets_xls(fn):
    b = xlrd.open_workbook(fn)
    return {re.sub(r'\s+', ' ', s.strip()): [b.sheet_by_name(s).row_values(i) for i in range(b.sheet_by_name(s).nrows)] for s in b.sheet_names()}

def find_col(rows, pat, upto=10):
    for r in rows[:upto]:
        for j, c in enumerate(r):
            if c and re.search(pat, str(c), re.I):
                return j

out = {}
# капитал/активы/качество/прибыль
for fn, rd in [('capital_assets.xlsx', sheets_xlsx), ('capital_assets_2025.xls', sheets_xls)]:
    for sh, rows in rd(fn).items():
        cE = find_col(rows, r'Собственный капитал по балансу'); cNI = find_col(rows, r'Превышение текущих')
        cA = find_col(rows, r'^Активы$'); cL = find_col(rows, r'Ссудный портфель'); cS3 = find_col(rows, r'3 стадии')
        cP = find_col(rows, r'Провизии'); c90 = find_col(rows, r'свыше 90')
        for r in rows:
            k = match(r[1] if len(r) > 1 else None)
            if not k: continue
            f = lambda j: (float(r[j]) if j is not None and r[j] not in (None, '') else None)
            out.setdefault(k, {}).setdefault(sh, {}).update(
                assets=f(cA), loans=f(cL), npl90=f(c90), npl90_sh=f(c90 + 1) if c90 is not None else None,
                s3=f(cS3), s3_sh=f(cS3 + 1) if cS3 is not None else None, prov=f(cP), equity=f(cE), ni_ytd=f(cNI))
# пруденциальные нормативы
for sh, rows in sheets_xlsx('prudential.xlsx').items():
    for r in rows:
        k = match(r[1] if len(r) > 1 else None)
        if k: out[k].setdefault(sh, {}).update(own_capital=r[2], k1=r[3], k1_2=r[4], k2=r[5])
# маржа и спред
for sh, rows in sheets_xlsx('nim_spread.xlsx').items():
    cM = find_col(rows, r'Процентная маржа'); cS = find_col(rows, r'Процентный спр')
    for r in rows:
        k = match(r[1] if len(r) > 1 else None)
        if k and cM is not None: out[k].setdefault(sh, {}).update(nim=r[cM], spread=r[cS])
json.dump(out, open('nbrk_monthly.json', 'w'), ensure_ascii=False, indent=1, default=str)
for k, v in out.items():
    print(k, sorted(v.keys()))
